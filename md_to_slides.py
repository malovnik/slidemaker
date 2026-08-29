#!/usr/bin/env python3
# -*- coding: utf-8 -*-

import argparse
import re
import os
import sys
from reportlab.lib import colors
from reportlab.lib.pagesizes import landscape, A4
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.enums import TA_CENTER
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, PageBreak, Flowable
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.lib.colors import HexColor, toColor, Color
from PIL import Image, ImageEnhance, ImageFilter, ImageOps
from reportlab.graphics.shapes import Drawing

# Поддерживаемые форматы фоновых изображений
SUPPORTED_BG_EXTENSIONS = ['.png', '.jpg', '.jpeg']

def find_background_image(background_path):
    """
    Ищет фоновое изображение с поддержкой разных форматов.

    Args:
        background_path: Путь к фоновому изображению (с расширением или без)

    Returns:
        str: Путь к найденному изображению или исходный путь если не найдено
    """
    # Если файл существует как есть - возвращаем его
    if os.path.exists(background_path):
        return background_path

    # Получаем базовое имя без расширения
    base_path = background_path.rsplit('.', 1)[0] if '.' in os.path.basename(background_path) else background_path

    # Ищем файл с разными расширениями
    for ext in SUPPORTED_BG_EXTENSIONS:
        candidate = f"{base_path}{ext}"
        if os.path.exists(candidate):
            print(f"Найдено фоновое изображение: {candidate}")
            return candidate

    # Если ничего не найдено, возвращаем исходный путь
    return background_path

# Определяем размеры страниц для разных форматов
FORMATS = {
    'horizontal': (1600, 900),   # 16:9
    'vertical': (900, 1600)      # 9:16
}

# Дефолтный размер (для совместимости)
PAGE_WIDTH = 1600
PAGE_HEIGHT = 900
PAGE_SIZE = (PAGE_WIDTH, PAGE_HEIGHT)

class GradientText(Flowable):
    """
    Flowable для отображения текста с градиентной заливкой
    """
    def __init__(self, text, font_name, font_size, start_color, end_color, width=None, height=None, alignment=TA_CENTER):
        Flowable.__init__(self)
        self.text = text
        self.font_name = font_name
        self.font_size = font_size
        self.start_color = start_color
        self.end_color = end_color
        self.width = width
        # Увеличиваем высоту для кириллических символов с выступающими элементами
        self.height = font_size * 1.5  # Увеличили коэффициент с 1.2 до 1.5
        self.alignment = alignment
        self.lines = [text]  # Изначально весь текст в одной строке
        print(f"GradientText: Создан с текстом '{text}', шрифтом '{font_name}', размером {font_size}")
        
    def wrap(self, availWidth, availHeight):
        """
        Определяет размер элемента и при необходимости делает перенос длинных строк
        """
        # Если ширина не задана, используем доступную ширину
        if self.width is None:
            self.width = availWidth
            
        # Настраиваем шрифт для измерения ширины текста
        try:
            # Это нужно для измерения текста в точных единицах
            from reportlab.pdfbase.pdfmetrics import stringWidth
            max_width = self.width * 0.95  # Оставляем небольшой запас
            
            # Разбиваем текст на строки, если он слишком длинный
            self.lines = []
            words = self.text.split()
            current_line = []
            current_width = 0
            
            for word in words:
                word_width = stringWidth(word, self.font_name, self.font_size)
                # Добавляем пробел, если строка не пустая
                space_width = stringWidth(' ', self.font_name, self.font_size) if current_line else 0
                
                # Если текущая строка + новое слово не превышают максимальную ширину
                if current_width + word_width + space_width <= max_width:
                    current_line.append(word)
                    current_width += word_width + space_width
                else:
                    # Если слово само по себе слишком длинное
                    if not current_line:
                        # Разбиваем длинное слово на части
                        start = 0
                        while start < len(word):
                            for end in range(min(len(word), start + 10), start, -1):
                                part = word[start:end]
                                part_width = stringWidth(part, self.font_name, self.font_size)
                                if part_width <= max_width:
                                    self.lines.append(part)
                                    start = end
                                    break
                            else:
                                # Если ни одна часть не подходит, добавляем самую маленькую
                                self.lines.append(word[start:start+1])
                                start += 1
                    else:
                        # Завершаем текущую строку и начинаем новую
                        self.lines.append(' '.join(current_line))
                        current_line = [word]
                        current_width = word_width
            
            # Добавляем последнюю строку, если она не пустая
            if current_line:
                self.lines.append(' '.join(current_line))
                
            # Если текст разбит на несколько строк, увеличиваем высоту
            if len(self.lines) > 1:
                self.height = self.font_size * 1.5 * len(self.lines)
        except ImportError:
            # Если не удалось импортировать stringWidth, оставляем текст как есть
            pass
        
        return (self.width, self.height)
        
    def draw(self):
        """
        Рисует текст с градиентной заливкой
        """
        # Настраиваем шрифт
        print(f"GradientText.draw: Устанавливаем шрифт '{self.font_name}', размер {self.font_size}")
        self.canv.setFont(self.font_name, self.font_size)
        
        # Рисуем каждую строку по отдельности
        y_position = self.height - self.font_size * 1.2  # Начальная позиция сверху
        
        for line in self.lines:
            print(f"GradientText.draw: Рисуем строку '{line}'")
            # Получаем ширину текста строки
            text_width = self.canv.stringWidth(line, self.font_name, self.font_size)
            print(f"GradientText.draw: Ширина текста: {text_width}")
            
            # Рисуем сначала текст целиком для определения правильной области
            self.canv.saveState()
            self.canv.setFillColorRGB(0, 0, 0, 0)  # Прозрачный цвет, только для определения области
            self.canv.drawString(0, 0, line)
            self.canv.restoreState()
            
            # Определяем количество шагов для градиента
            steps = 50
            
            # Вычисляем шаг для изменения цвета
            r_step = (self.end_color.red - self.start_color.red) / steps
            g_step = (self.end_color.green - self.start_color.green) / steps
            b_step = (self.end_color.blue - self.start_color.blue) / steps
            
            # Вычисляем смещение для центрирования текста
            if self.alignment == TA_CENTER:
                x_offset = (self.width - text_width) / 2
            else:
                x_offset = 0
                
            # Позиционируем текст по вертикали, чтобы учесть нижние выступающие элементы
            y_offset = y_position + self.font_size * 0.3  # Подъем по вертикали для учета нижних выступов букв
                
            # Определяем ширину каждого сегмента
            segment_width = text_width / steps
            
            # Рисуем каждый сегмент с соответствующим цветом
            for i in range(steps):
                # Вычисляем текущий цвет
                current_r = self.start_color.red + (r_step * i)
                current_g = self.start_color.green + (g_step * i)
                current_b = self.start_color.blue + (b_step * i)
                
                current_color = Color(current_r, current_g, current_b)
                
                # Устанавливаем цвет
                self.canv.setFillColor(current_color)
                
                # Определяем координаты для текущего сегмента
                x1 = x_offset + (i * segment_width)
                
                # Рисуем сегмент текста
                self.canv.saveState()
                path = self.canv.beginPath()
                path.rect(x1, 0, segment_width, self.height)
                self.canv.clipPath(path, stroke=0)
                self.canv.drawString(x_offset, y_offset, line)  # Рисуем текущую строку
                self.canv.restoreState()
                
            # Уменьшаем y_position для следующей строки
            y_position -= self.font_size * 1.2

    def register_fonts(self):
        """
        Регистрирует шрифты для использования в PDF
        """
        try:
            # Проверяем наличие шрифта напрямую
            direct_font_path = os.path.join("fonts", f"{self.font_name}.ttf")
            print(f"Проверяем наличие шрифта: {direct_font_path}")
            print(f"Текущая директория: {os.getcwd()}")
            print(f"Файл существует: {os.path.exists(direct_font_path)}")
            
            if os.path.exists(direct_font_path):
                print(f"Регистрируем шрифт: {direct_font_path}")
                pdfmetrics.registerFont(TTFont(self.font_name, direct_font_path))
                print(f"Шрифт успешно зарегистрирован: {self.font_name}")
                return True
            
            # Проверяем наличие шрифта с суффиксами Regular и Bold
            regular_font_path = os.path.join("fonts", f"{self.font_name}-Regular.ttf")
            bold_font_path = os.path.join("fonts", f"{self.font_name}-Bold.ttf")
            
            print(f"Проверяем наличие шрифта Regular: {regular_font_path}")
            print(f"Файл существует: {os.path.exists(regular_font_path)}")
            print(f"Проверяем наличие шрифта Bold: {bold_font_path}")
            print(f"Файл существует: {os.path.exists(bold_font_path)}")
            
            if os.path.exists(regular_font_path) and os.path.exists(bold_font_path):
                print(f"Регистрируем шрифты: {regular_font_path}, {bold_font_path}")
                pdfmetrics.registerFont(TTFont(f"{self.font_name}", regular_font_path))
                pdfmetrics.registerFont(TTFont(f"{self.font_name}-Bold", bold_font_path))
                print(f"Шрифты успешно зарегистрированы: {self.font_name}, {self.font_name}-Bold")
                return True
            
            # Если шрифты не найдены, используем стандартные
            print("Шрифты не найдены, используем стандартные")
            return False
        except Exception as e:
            print(f"Ошибка при регистрации шрифтов: {e}")
            return False

class MarkdownSlideConverter:
    def __init__(self, font_path=None, font_name="Montserrat", text_font_name=None,
                 text_font_path=None,
                 background_image="bg",
                 heading_color="1b1b1b", text_color="1b1b1b",
                 accent_color="ff6b6b",
                 heading_gradient="no", heading_gradient_end="1b1b1b",
                 text_gradient="no", text_gradient_end="1b1b1b",
                 bg_effect="none", bg_effect_intensity="0",
                 format="horizontal",
                 heading_font_size=50, heading_width=90,
                 text_font_size=36, text_width=85):
        """
        Инициализирует конвертер слайдов из Markdown.

        Args:
            font_path: Путь к папке с шрифтами заголовков
            font_name: Имя шрифта для заголовков
            text_font_name: Имя шрифта для текста (если None, используется font_name)
            text_font_path: Путь к папке со шрифтами текста (если None, используется font_path)
            background_image: Путь к фоновому изображению (поддерживаются png, jpg, jpeg)
            heading_color: Цвет заголовков в HEX-формате без #
            text_color: Цвет текста в HEX-формате без #
            accent_color: Цвет акцента для **bold** текста в HEX-формате без #
            heading_gradient: Использовать градиент для заголовков (yes/no)
            heading_gradient_end: Конечный цвет градиента заголовков
            text_gradient: Использовать градиент для текста (yes/no)
            text_gradient_end: Конечный цвет градиента текста
            bg_effect: Эффект для фона (lighten/darken/none)
            bg_effect_intensity: Интенсивность эффекта (1-10)
            format: Формат презентации (horizontal/vertical)
            heading_font_size: Размер шрифта заголовков (15-200)
            heading_width: Ширина заголовков в процентах (30-120)
            text_font_size: Размер шрифта текста (12-80)
            text_width: Ширина текста в процентах (30-120)
        """
        self.font_name = font_name
        self.text_font_name = text_font_name if text_font_name else font_name
        self.font_path = font_path
        self.text_font_path = text_font_path if text_font_path else font_path

        # Устанавливаем размер страницы в зависимости от формата
        self.format = format
        self.page_width, self.page_height = FORMATS.get(format, FORMATS['horizontal'])
        self.page_size = (self.page_width, self.page_height)
        
        # Устанавливаем размеры шрифтов и ширину (с ограничениями)
        self.heading_font_size = max(15, min(200, int(heading_font_size)))
        self.heading_width = max(30, min(120, int(heading_width)))
        self.text_font_size = max(12, min(80, int(text_font_size)))
        self.text_width = max(30, min(120, int(text_width)))
        
        # Ищем фоновое изображение с поддержкой разных форматов
        self.background_image = find_background_image(background_image)
        self.heading_gradient = heading_gradient
        self.text_gradient = text_gradient
        self.bg_effect = bg_effect
        self.bg_effect_intensity = int(bg_effect_intensity)
        
        # Обработка фона (кадрирование + эффекты)
        self.processed_bg = None
        if os.path.exists(self.background_image):
            self.process_background()
        
        # Проверяем наличие фонового изображения
        self.has_background = os.path.exists(self.background_image)
        if not self.has_background:
            print(f"Предупреждение: Фоновое изображение не найдено (искали: {background_image} с расширениями {', '.join(SUPPORTED_BG_EXTENSIONS)}). Будет использован белый фон.")
        
        # Цвет заголовков
        try:
            self.heading_color = HexColor(f'#{heading_color}')
            self.heading_color_hex = heading_color  # Сохраняем hex строку для PPTX

            # Если используется градиент, устанавливаем конечный цвет
            if self.heading_gradient == "yes":
                self.heading_gradient_end_color = HexColor(f'#{heading_gradient_end}')
                print(f"Используется градиент для заголовков: от #{heading_color} до #{heading_gradient_end}")
            else:
                print(f"Используется цвет заголовков: #{heading_color}")
        except Exception as e:
            print(f"Ошибка при установке цвета заголовков: {e}")
            print("Используется цвет заголовков по умолчанию (#1b1b1b)")
            self.heading_color = HexColor('#1b1b1b')
            self.heading_color_hex = '1b1b1b'
            self.heading_gradient = "no"

        # Цвет текста
        try:
            self.text_color = HexColor(f'#{text_color}')
            self.text_color_hex = text_color  # Сохраняем hex строку для PPTX

            # Если используется градиент, устанавливаем конечный цвет
            if self.text_gradient == "yes":
                self.text_gradient_end_color = HexColor(f'#{text_gradient_end}')
                print(f"Используется градиент для текста: от #{text_color} до #{text_gradient_end}")
            else:
                print(f"Используется цвет текста: #{text_color}")
        except Exception as e:
            print(f"Ошибка при установке цвета текста: {e}")
            print("Используется цвет текста по умолчанию (#1b1b1b)")
            self.text_color = HexColor('#1b1b1b')
            self.text_color_hex = '1b1b1b'
            self.text_gradient = "no"

        # Акцентный цвет для **bold** текста
        try:
            self.accent_color = HexColor(f'#{accent_color}')
            self.accent_color_hex = accent_color
            print(f"Используется акцентный цвет: #{accent_color}")
        except Exception as e:
            print(f"Ошибка при установке акцентного цвета: {e}")
            self.accent_color = HexColor('#ff6b6b')
            self.accent_color_hex = 'ff6b6b'

        # Регистрируем шрифты
        self.register_fonts()
        
        # Создаем стили для разных уровней заголовков и текста
        self.create_styles()
    
    def crop_to_aspect_ratio(self, img, target_width, target_height):
        """
        Обрезает изображение до нужного соотношения сторон (crop по центру).

        Args:
            img: PIL Image объект
            target_width: Целевая ширина (для расчёта соотношения)
            target_height: Целевая высота (для расчёта соотношения)

        Returns:
            PIL Image с обрезанным изображением
        """
        img_width, img_height = img.size
        target_ratio = target_width / target_height
        img_ratio = img_width / img_height

        if img_ratio > target_ratio:
            # Изображение шире - обрезаем по бокам
            new_width = int(img_height * target_ratio)
            left = (img_width - new_width) // 2
            img = img.crop((left, 0, left + new_width, img_height))
        elif img_ratio < target_ratio:
            # Изображение выше - обрезаем сверху и снизу
            new_height = int(img_width / target_ratio)
            top = (img_height - new_height) // 2
            img = img.crop((0, top, img_width, top + new_height))

        return img

    def process_background(self):
        """Обрабатывает фоновое изображение: кадрирование и эффекты"""
        try:
            # Загружаем изображение
            img = Image.open(self.background_image)

            # Исправляем EXIF ориентацию (для фото с телефонов)
            img = ImageOps.exif_transpose(img)

            # Преобразуем в RGBA для поддержки прозрачности
            if img.mode != 'RGBA':
                img = img.convert('RGBA')

            # Кадрируем до нужного соотношения сторон
            img = self.crop_to_aspect_ratio(img, self.page_width, self.page_height)
            print(f"Изображение кадрировано до соотношения {self.page_width}:{self.page_height}")

            # Применяем эффект осветления
            if self.bg_effect == "lighten":
                alpha = int(255 * (self.bg_effect_intensity / 10))
                white_layer = Image.new('RGBA', img.size, (255, 255, 255, alpha))
                img = Image.alpha_composite(img, white_layer)
                print(f"Применён эффект осветления с интенсивностью {self.bg_effect_intensity}/10")

            # Применяем эффект затемнения
            elif self.bg_effect == "darken":
                alpha = int(255 * (self.bg_effect_intensity / 10))
                black_layer = Image.new('RGBA', img.size, (0, 0, 0, alpha))
                img = Image.alpha_composite(img, black_layer)
                print(f"Применён эффект затемнения с интенсивностью {self.bg_effect_intensity}/10")

            # Сохраняем обработанное изображение во временный файл
            processed_bg_path = f"{self.background_image.rsplit('.', 1)[0]}_processed.png"
            img.save(processed_bg_path)
            self.processed_bg = processed_bg_path
            print(f"Обработанное фоновое изображение сохранено как: {processed_bg_path}")

        except Exception as e:
            print(f"Ошибка при обработке фонового изображения: {e}")
            print("Будет использовано оригинальное изображение")
            self.processed_bg = None
    
    def create_styles(self):
        """Создает стили для разных элементов презентации"""
        self.styles = getSampleStyleSheet()
        
        # Вычисляем пропорциональные размеры для разных стилей заголовков
        title_size = self.heading_font_size
        subtitle_size = int(self.heading_font_size * 0.72)  # 36/50
        section_size = int(self.heading_font_size * 0.92)   # 46/50
        sub_size = int(self.heading_font_size * 0.84)       # 42/50
        content_size = self.text_font_size
        
        # Вычисляем отступы для ширины (базовая ширина = page_width - 160)
        # Формула: indent = base_width * (100 - width) / 100 / 2
        base_width = self.page_width - 160  # Доступная ширина при стандартных margins
        heading_indent = int(base_width * (100 - self.heading_width) / 100 / 2)
        text_indent = int(base_width * (100 - self.text_width) / 100 / 2)
        
        # Устанавливаем стили с учетом градиентов
        self.styles.add(ParagraphStyle(
            name='TitleSlide',
            fontName=self.font_bold,
            fontSize=title_size,
            alignment=TA_CENTER,
            spaceAfter=int(title_size * 0.8),
            leading=int(title_size * 1.2),
            textColor=self.heading_color,
            leftIndent=heading_indent,
            rightIndent=heading_indent
        ))
        self.styles.add(ParagraphStyle(
            name='SubTitleSlide',
            fontName=self.font_regular,
            fontSize=subtitle_size,
            alignment=TA_CENTER,
            spaceAfter=int(subtitle_size * 1.4),
            leading=int(subtitle_size * 1.22),
            textColor=self.heading_color,
            leftIndent=heading_indent,
            rightIndent=heading_indent
        ))
        self.styles.add(ParagraphStyle(
            name='SectionSlide',
            fontName=self.font_bold,
            fontSize=section_size,
            alignment=TA_CENTER,
            spaceAfter=int(section_size * 0.65),
            leading=int(section_size * 1.13),
            textColor=self.heading_color,
            leftIndent=heading_indent,
            rightIndent=heading_indent
        ))
        self.styles.add(ParagraphStyle(
            name='SubSlide',
            fontName=self.font_bold,
            fontSize=sub_size,
            alignment=TA_CENTER,
            spaceAfter=int(sub_size * 0.6),
            leading=int(sub_size * 1.14),
            textColor=self.heading_color,
            leftIndent=heading_indent,
            rightIndent=heading_indent
        ))
        self.styles.add(ParagraphStyle(
            name='Content',
            fontName=self.text_font_regular,
            fontSize=content_size,
            alignment=TA_CENTER,
            spaceAfter=int(content_size * 0.42),
            leading=int(content_size * 1.17),
            textColor=self.text_color,
            leftIndent=text_indent,
            rightIndent=text_indent
        ))
    
    def parse_inline_markdown(self, text, font_name):
        """
        Парсит inline markdown форматирование и возвращает HTML для Paragraph.
        **bold** -> акцентный цвет
        *italic* -> убираем звёздочки

        Args:
            text: Исходный текст с markdown
            font_name: Имя шрифта для сохранения в тегах

        Returns:
            str: Текст с ReportLab XML-тегами для стилизации
        """
        import re

        # Обрабатываем **bold** -> только меняем цвет (шрифт сохраняется внешним тегом)
        text = re.sub(
            r'\*\*(.+?)\*\*',
            f'<font color="#{self.accent_color_hex}">\\1</font>',
            text
        )

        # Обрабатываем *italic* -> убираем звёздочки
        text = re.sub(r'\*(.+?)\*', r'\1', text)

        # ВАЖНО: Оборачиваем ВЕСЬ текст в тег с правильным шрифтом
        # Это гарантирует, что после закрывающих тегов </font> шрифт не сбросится
        return f'<font face="{font_name}">{text}</font>'

    def create_text_element(self, text, style_name, is_heading=True):
        """
        Создает элемент текста с учетом градиента и markdown форматирования.

        Args:
            text: Текст для отображения
            style_name: Имя стиля из self.styles
            is_heading: Является ли текст заголовком (для выбора цвета)

        Returns:
            Flowable: Элемент для добавления в PDF
        """
        style = self.styles[style_name]

        # Определяем, использовать ли градиент
        use_gradient = False
        if is_heading and self.heading_gradient == "yes":
            use_gradient = True
            start_color = self.heading_color
            end_color = self.heading_gradient_end_color
        elif not is_heading and self.text_gradient == "yes":
            use_gradient = True
            start_color = self.text_color
            end_color = self.text_gradient_end_color

        print(f"create_text_element: Текст '{text}', стиль '{style_name}', заголовок: {is_heading}, градиент: {use_gradient}")
        print(f"create_text_element: Шрифт '{style.fontName}', размер {style.fontSize}")

        # Для контента (не заголовка) парсим markdown форматирование
        if not is_heading:
            # Передаём имя шрифта для сохранения в тегах
            parsed_text = self.parse_inline_markdown(text, style.fontName)

            # Если используем градиент, убираем HTML теги (градиент не поддерживает их)
            if use_gradient:
                # Убираем теги, оставляем только текст
                import re
                clean_text = re.sub(r'<[^>]+>', '', parsed_text)
                return GradientText(
                    text=clean_text,
                    font_name=style.fontName,
                    font_size=style.fontSize,
                    start_color=start_color,
                    end_color=end_color,
                    alignment=style.alignment
                )
            else:
                # Paragraph поддерживает HTML-подобные теги
                return Paragraph(parsed_text, style)

        # Для заголовков - тоже парсим markdown (для поддержки **акцента**)
        parsed_text = self.parse_inline_markdown(text, style.fontName)

        if use_gradient:
            # Для градиента убираем HTML теги
            import re
            clean_text = re.sub(r'<[^>]+>', '', parsed_text)
            return GradientText(
                text=clean_text,
                font_name=style.fontName,
                font_size=style.fontSize,
                start_color=start_color,
                end_color=end_color,
                alignment=style.alignment
            )
        else:
            return Paragraph(parsed_text, style)
    
    def _register_single_font(self, font_name, prefix="", custom_font_path=None):
        """
        Регистрирует один шрифт и возвращает имена regular и bold вариантов.

        Args:
            font_name: Имя шрифта
            prefix: Префикс для имени (для различения шрифтов заголовка и текста)
            custom_font_path: Путь к папке со шрифтом (если None, используется self.font_path)

        Returns:
            Tuple (font_regular, font_bold)
        """
        # Используем переданный путь или путь по умолчанию
        font_dir = custom_font_path or self.font_path
        base_dir = font_dir if font_dir else "fonts"

        # Функция для поиска файла шрифта с разными расширениями
        def find_font_file(base_path):
            """Ищет файл шрифта с расширением .ttf или .otf"""
            for ext in ['.ttf', '.otf']:
                path = base_path + ext
                if os.path.exists(path):
                    return path
            return None

        # Пути без расширения
        direct_base = os.path.join(base_dir, font_name)
        regular_base = os.path.join(base_dir, f"{font_name}-Regular")
        bold_base = os.path.join(base_dir, f"{font_name}-Bold")

        # Ищем файлы шрифтов
        direct_font_path = find_font_file(direct_base)
        regular_font_path = find_font_file(regular_base)
        bold_font_path = find_font_file(bold_base)

        # Проверяем, существует ли файл шрифта без суффикса
        if direct_font_path:
            font_regular = f"{prefix}{font_name}"
            font_bold = f"{prefix}{font_name}"

            # Регистрируем только если ещё не зарегистрирован
            try:
                pdfmetrics.getFont(font_regular)
                print(f"[DEBUG] Шрифт {font_regular} уже зарегистрирован")
            except KeyError:
                pdfmetrics.registerFont(TTFont(font_regular, direct_font_path))
                print(f"[DEBUG] Шрифт {font_name} загружен из файла: {direct_font_path}")

            return font_regular, font_bold

        # Проверяем варианты с суффиксами
        if regular_font_path and bold_font_path:
            font_regular = f"{prefix}{font_name}-Regular"
            font_bold = f"{prefix}{font_name}-Bold"

            # Регистрируем Regular если не зарегистрирован
            try:
                pdfmetrics.getFont(font_regular)
            except KeyError:
                pdfmetrics.registerFont(TTFont(font_regular, regular_font_path))
                print(f"Шрифт {font_regular} загружен из {regular_font_path}")

            # Регистрируем Bold если не зарегистрирован (ОТДЕЛЬНО!)
            try:
                pdfmetrics.getFont(font_bold)
            except KeyError:
                pdfmetrics.registerFont(TTFont(font_bold, bold_font_path))
                print(f"Шрифт {font_bold} загружен из {bold_font_path}")

            return font_regular, font_bold

        # Если есть только Regular вариант
        if regular_font_path:
            font_regular = f"{prefix}{font_name}-Regular"
            font_bold = font_regular  # Используем Regular вместо Bold

            try:
                pdfmetrics.getFont(font_regular)
            except KeyError:
                pdfmetrics.registerFont(TTFont(font_regular, regular_font_path))
                print(f"Шрифт {font_name} (только Regular) загружен из {regular_font_path}")

            return font_regular, font_bold

        # Fallback
        print(f"Шрифт {font_name} не найден в {base_dir}, используется Helvetica")
        return "Helvetica", "Helvetica-Bold"

    def register_fonts(self):
        """Регистрирует шрифты для использования в PDF"""
        try:
            print(f"Текущая директория: {os.getcwd()}")
            print(f"Шрифт заголовков: {self.font_name}, путь: {self.font_path}")
            print(f"Шрифт текста: {self.text_font_name}, путь: {self.text_font_path}")

            # Регистрируем шрифт заголовков
            self.font_regular, self.font_bold = self._register_single_font(
                self.font_name, custom_font_path=self.font_path
            )

            # Регистрируем шрифт текста (если отличается от заголовков)
            if self.text_font_name != self.font_name:
                self.text_font_regular, self.text_font_bold = self._register_single_font(
                    self.text_font_name, prefix="text_", custom_font_path=self.text_font_path
                )
            else:
                self.text_font_regular = self.font_regular
                self.text_font_bold = self.font_bold

        except Exception as e:
            print(f"Ошибка при загрузке шрифтов: {e}")
            self.font_regular = "Helvetica"
            self.font_bold = "Helvetica-Bold"
            self.text_font_regular = "Helvetica"
            self.text_font_bold = "Helvetica-Bold"
            print("Используются стандартные шрифты Helvetica")
    
    def parse_markdown(self, md_file):
        """
        Разбирает Markdown-файл и создает структуру слайдов.
        
        Args:
            md_file: Путь к Markdown-файлу
            
        Returns:
            list: Список слайдов с их содержимым
        """
        slides = []
        current_slide = None
        
        # Сначала читаем весь файл
        with open(md_file, 'r', encoding='utf-8') as f:
            content = f.readlines()
        
        # Анализируем структуру документа для определения иерархии заголовков
        header_levels = {}
        for line in content:
            line = line.strip()
            if not line:
                continue
            
            # Ищем заголовки (строки, начинающиеся с # ## ### и т.д.)
            header_match = re.match(r'^(#{1,6})\s+(.+)$', line)
            if header_match:
                level = len(header_match.group(1))  # Количество символов #
                if level not in header_levels:
                    header_levels[level] = 0
                header_levels[level] += 1
        
        # Определяем уровни заголовков для разных типов слайдов
        print(f"Обнаружены уровни заголовков: {sorted(header_levels.keys())}")
        
        # Если есть только один уровень заголовков, используем его для основных слайдов
        if len(header_levels) == 1:
            main_level = list(header_levels.keys())[0]
            print(f"Обнаружен один уровень заголовков ({main_level}). Каждый заголовок будет отдельным слайдом.")
            slide_level = main_level
            section_level = None
            subsection_level = None
        # Если есть несколько уровней, используем их для иерархии
        else:
            levels = sorted(header_levels.keys())
            print(f"Обнаружена иерархия заголовков: {levels}")
            slide_level = levels[0]  # Самый верхний уровень для основных слайдов
            section_level = levels[1] if len(levels) > 1 else None
            subsection_level = levels[2] if len(levels) > 2 else None
        
        # Теперь обрабатываем содержимое с учётом определённой иерархии
        for line in content:
            line = line.strip()
            if not line:
                continue
            
            # Проверяем, является ли строка заголовком
            header_match = re.match(r'^(#{1,6})\s+(.+)$', line)
            
            if header_match:
                level = len(header_match.group(1))
                title = header_match.group(2).strip()
                
                # Основной слайд
                if level == slide_level:
                    # Если это первый слайд, считаем его титульным
                    if not slides:
                        current_slide = {'type': 'title', 'title': title, 'subtitle': None, 'content': []}
                    else:
                        current_slide = {'type': 'section', 'title': title, 'content': []}
                    slides.append(current_slide)
                
                # Раздел (если есть второй уровень иерархии)
                elif section_level and level == section_level:
                    current_slide = {'type': 'section', 'title': title, 'content': []}
                    slides.append(current_slide)
                
                # Подраздел (если есть третий уровень иерархии)
                elif subsection_level and level == subsection_level:
                    current_slide = {'type': 'subsection', 'title': title, 'content': []}
                    slides.append(current_slide)
            
            # Добавляем содержимое (элементы списка)
            elif line.startswith('- ') and current_slide:
                content = line[2:].strip()
                current_slide['content'].append(content)
        
        # Если в титульном слайде есть контент, попробуем выделить подзаголовок
        if slides and slides[0]['type'] == 'title' and slides[0]['content']:
            slides[0]['subtitle'] = slides[0]['content'][0]
            slides[0]['content'] = slides[0]['content'][1:]
        
        return slides
    
    def add_background(self, canvas, doc):
        """
        Добавляет фоновое изображение на каждую страницу PDF.

        Args:
            canvas: Канва страницы
            doc: Документ PDF
        """
        if self.has_background:
            try:
                # Получаем путь к изображению (обработанному или оригинальному)
                bg_image_path = self.processed_bg if self.processed_bg else self.background_image

                # Получаем размеры страницы
                page_width, page_height = self.page_size
                
                # Загружаем и масштабируем изображение на всю страницу
                img = Image.open(bg_image_path)
                img_width, img_height = img.size
                
                # Вычисляем коэффициент масштабирования для заполнения всей страницы
                scale_width = page_width / img_width
                scale_height = page_height / img_height
                scale = max(scale_width, scale_height)
                
                # Вычисляем новые размеры изображения
                new_width = img_width * scale
                new_height = img_height * scale
                
                # Вычисляем координаты для центрирования изображения
                x_offset = (page_width - new_width) / 2
                y_offset = (page_height - new_height) / 2
                
                # Рисуем изображение на канве
                canvas.drawImage(bg_image_path, x_offset, y_offset, width=new_width, height=new_height)
            except Exception as e:
                print(f"Ошибка при добавлении фона: {e}")
    
    def create_pdf(self, slides, output_file):
        """
        Создает PDF-презентацию из структуры слайдов.

        Args:
            slides: Список слайдов
            output_file: Путь к выходному PDF-файлу
        """
        # Создаем документ с обработчиком фона
        doc = SimpleDocTemplate(
            output_file,
            pagesize=self.page_size,
            rightMargin=80,
            leftMargin=80,
            topMargin=80,
            bottomMargin=100  # Больше отступ снизу для возможности размещения видео
        )
        
        # Создаем список элементов для PDF
        all_stories = []
        
        for i, slide in enumerate(slides):
            story = []
            elements = []
            
            # Обрабатываем титульный слайд специальным образом
            if slide['type'] == 'title':
                elements.append(self.create_text_element(slide['title'], 'TitleSlide', is_heading=True))
                
                # Добавляем подзаголовок если он есть
                if 'subtitle' in slide and slide['subtitle']:
                    elements.append(Spacer(1, 60))  # Увеличиваем отступ между заголовком и подзаголовком
                    elements.append(self.create_text_element(slide['subtitle'], 'SubTitleSlide', is_heading=True))
            
            # Другие типы слайдов
            elif slide['type'] == 'section':
                elements.append(self.create_text_element(slide['title'], 'SectionSlide', is_heading=True))
            elif slide['type'] == 'subsection':
                elements.append(self.create_text_element(slide['title'], 'SubSlide', is_heading=True))
            
            # Добавляем содержимое слайда
            if slide['content']:
                elements.append(Spacer(1, 50))  # Увеличенный отступ после заголовка
                
                for content_item in slide['content']:
                    elements.append(self.create_text_element(content_item, 'Content', is_heading=False))
                    elements.append(Spacer(1, 25))  # Увеличенный отступ между пунктами
            
            # Оцениваем общую высоту содержимого
            total_height = 0
            for element in elements:
                if isinstance(element, Paragraph) or isinstance(element, GradientText):
                    # Оцениваем высоту
                    w, h = element.wrap(self.page_width - 160, self.page_height)  # Учитываем увеличенные поля
                    total_height += h
                elif isinstance(element, Spacer):
                    total_height += element.height
            
            # Вычисляем отступ сверху для вертикального центрирования
            available_height = self.page_height - 180  # Учитываем увеличенные отступы сверху и снизу
            top_padding = max(0, (available_height - total_height) / 2)
            
            # Добавляем верхний отступ для вертикального центрирования
            story.append(Spacer(1, top_padding))
            
            # Добавляем все элементы на текущий слайд
            for element in elements:
                story.append(element)
            
            # Добавляем разрыв страницы в конце слайда, кроме последнего
            if i < len(slides) - 1:
                story.append(PageBreak())
            
            # Добавляем содержимое слайда в общий список
            all_stories.extend(story)
        
        # Строим документ с фоном
        doc.build(all_stories, onFirstPage=self.add_background, onLaterPages=self.add_background)
        print(f"Презентация сохранена в {output_file}")
        
        # Удаляем временные файлы
        if self.processed_bg and os.path.exists(self.processed_bg):
            try:
                os.remove(self.processed_bg)
                print(f"Временное обработанное изображение удалено")
            except:
                pass

    def create_pptx(self, slides, output_file):
        """
        Создает PPTX-презентацию из структуры слайдов.

        Args:
            slides: Список слайдов
            output_file: Путь к выходному PPTX-файлу
        """
        try:
            from pptx import Presentation
            from pptx.util import Inches, Pt
            from pptx.dml.color import RGBColor
            from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
        except ImportError as e:
            import traceback
            traceback.print_exc()
            raise ImportError(f"Ошибка импорта python-pptx: {str(e)}. Убедитесь, что библиотека установлена: pip install python-pptx")

        # Создаём презентацию с нужным размером
        prs = Presentation()

        # Устанавливаем размер слайдов
        if self.format == 'vertical':
            prs.slide_width = Inches(9)
            prs.slide_height = Inches(16)
        else:
            prs.slide_width = Inches(16)
            prs.slide_height = Inches(9)

        # Используем пустой layout
        blank_layout = prs.slide_layouts[6]

        # Определяем имена шрифтов для PPTX (убираем суффиксы)
        def get_pptx_font_name(font_name):
            for suffix in ['-Regular', '-Bold', '-Italic', '-Light', '-Medium',
                          '-SemiBold', '-ExtraBold', '-Black', '-Thin', '-ExtraLight',
                          '-LightItalic', '-MediumItalic', '-BoldItalic', '-BlackItalic']:
                if font_name.endswith(suffix):
                    return font_name[:-len(suffix)]
            return font_name

        heading_font = get_pptx_font_name(self.font_name)
        text_font = get_pptx_font_name(self.text_font_name or self.font_name)

        for slide_data in slides:
            slide = prs.slides.add_slide(blank_layout)

            # Добавляем фоновое изображение
            bg_path = self.processed_bg if self.processed_bg else self.background_image
            if os.path.exists(bg_path):
                slide.shapes.add_picture(
                    bg_path, 0, 0,
                    width=prs.slide_width,
                    height=prs.slide_height
                )

            # Заголовок
            title = slide_data.get('title', '')
            if title:
                # Рассчитываем позицию для центрирования
                title_top = prs.slide_height * 0.3
                title_height = Inches(1.5)

                # Вычисляем отступы для ширины заголовка
                heading_margin = prs.slide_width * (100 - self.heading_width) / 100 / 2
                
                title_box = slide.shapes.add_textbox(
                    heading_margin, title_top,
                    prs.slide_width - heading_margin * 2, title_height
                )
                title_frame = title_box.text_frame
                title_frame.word_wrap = True

                p = title_frame.paragraphs[0]
                p.text = title
                p.font.size = Pt(int(self.heading_font_size * 0.88))  # Пропорциональный размер для PPTX
                p.font.bold = True
                p.font.name = heading_font
                p.alignment = PP_ALIGN.CENTER

                # Устанавливаем цвет заголовка
                heading_hex = self.heading_color_hex
                p.font.color.rgb = RGBColor(
                    int(heading_hex[0:2], 16),
                    int(heading_hex[2:4], 16),
                    int(heading_hex[4:6], 16)
                )

            # Контент
            content = slide_data.get('content', [])
            if content:
                content_top = prs.slide_height * 0.5
                content_height = prs.slide_height * 0.4

                # Вычисляем отступы для ширины текста
                text_margin = prs.slide_width * (100 - self.text_width) / 100 / 2
                
                content_box = slide.shapes.add_textbox(
                    text_margin, content_top,
                    prs.slide_width - text_margin * 2, content_height
                )
                content_frame = content_box.text_frame
                content_frame.word_wrap = True

                text_hex = self.text_color_hex
                text_rgb = RGBColor(
                    int(text_hex[0:2], 16),
                    int(text_hex[2:4], 16),
                    int(text_hex[4:6], 16)
                )

                # Акцентный цвет для PPTX
                accent_hex = self.accent_color_hex
                accent_rgb = RGBColor(
                    int(accent_hex[0:2], 16),
                    int(accent_hex[2:4], 16),
                    int(accent_hex[4:6], 16)
                )

                for i, item in enumerate(content):
                    if i == 0:
                        p = content_frame.paragraphs[0]
                    else:
                        p = content_frame.add_paragraph()

                    p.alignment = PP_ALIGN.CENTER
                    p.space_after = Pt(12)

                    # Парсим markdown и создаём runs с правильными цветами
                    import re
                    # Находим все части текста: обычный текст и **акцент**
                    pattern = r'(\*\*(.+?)\*\*)'
                    last_end = 0
                    has_runs = False

                    for match in re.finditer(r'\*\*(.+?)\*\*', item):
                        has_runs = True
                        # Размер шрифта для текста (пропорциональный)
                        pptx_text_size = Pt(int(self.text_font_size * 0.78))  # Пропорциональный размер для PPTX
                        
                        # Добавляем текст до акцента
                        if match.start() > last_end:
                            run = p.add_run()
                            run.text = item[last_end:match.start()]
                            run.font.size = pptx_text_size
                            run.font.color.rgb = text_rgb
                            run.font.name = text_font

                        # Добавляем акцентный текст
                        run = p.add_run()
                        run.text = match.group(1)
                        run.font.size = pptx_text_size
                        run.font.color.rgb = accent_rgb
                        run.font.name = text_font
                        run.font.bold = True

                        last_end = match.end()

                    # Размер шрифта для текста (пропорциональный)
                    pptx_text_size = Pt(int(self.text_font_size * 0.78))
                    
                    # Добавляем оставшийся текст после последнего акцента
                    if has_runs:
                        if last_end < len(item):
                            run = p.add_run()
                            run.text = item[last_end:]
                            run.font.size = pptx_text_size
                            run.font.color.rgb = text_rgb
                            run.font.name = text_font
                    else:
                        # Нет акцентов — просто добавляем текст
                        # Убираем *italic* если есть
                        clean_item = re.sub(r'\*(.+?)\*', r'\1', item)
                        run = p.add_run()
                        run.text = clean_item
                        run.font.size = pptx_text_size
                        run.font.color.rgb = text_rgb
                        run.font.name = text_font

        prs.save(output_file)
        print(f"PPTX презентация сохранена в {output_file}")

    def create_jpg_archive(self, slides, output_file, pdf_path=None):
        """
        Создаёт ZIP-архив с JPG изображениями слайдов.

        Args:
            slides: Список слайдов
            output_file: Путь к выходному ZIP-файлу
            pdf_path: Путь к PDF файлу (если уже создан)
        """
        import tempfile
        import zipfile

        try:
            from pdf2image import convert_from_path
        except ImportError:
            raise ImportError("Библиотека pdf2image не установлена")

        # Если PDF не передан, создаём временный
        if not pdf_path:
            with tempfile.NamedTemporaryFile(suffix='.pdf', delete=False) as tmp:
                temp_pdf = tmp.name
            self.create_pdf(slides, temp_pdf)
            pdf_path = temp_pdf
            cleanup_pdf = True
        else:
            cleanup_pdf = False

        try:
            # Определяем DPI для нужного разрешения
            # Формула: target_pixels = dpi * (page_points / 72)
            # Значит: dpi = target_pixels * 72 / page_points
            if self.format == 'vertical':
                # 1080x1920 для вертикального
                target_width = 1080
                dpi = int(target_width * 72 / self.page_width)
            else:
                # 1920x1080 для горизонтального
                target_width = 1920
                dpi = int(target_width * 72 / self.page_width)

            # Конвертируем PDF в изображения
            images = convert_from_path(pdf_path, dpi=dpi)

            # Создаём ZIP архив
            with tempfile.TemporaryDirectory() as tmpdir:
                image_paths = []

                for i, image in enumerate(images):
                    img_path = os.path.join(tmpdir, f'slide_{i+1:03d}.jpg')
                    image.save(img_path, 'JPEG', quality=95)
                    image_paths.append(img_path)

                # Создаём архив
                with zipfile.ZipFile(output_file, 'w', zipfile.ZIP_DEFLATED) as zipf:
                    for img_path in image_paths:
                        zipf.write(img_path, os.path.basename(img_path))

            print(f"JPG архив сохранён в {output_file}")

        finally:
            if cleanup_pdf and os.path.exists(pdf_path):
                try:
                    os.remove(pdf_path)
                except:
                    pass


def main():
    parser = argparse.ArgumentParser(description='Конвертер Markdown в PDF-презентацию')
    parser.add_argument('input_file', help='Входной Markdown-файл')
    parser.add_argument('-o', '--output', help='Выходной PDF-файл')
    parser.add_argument('-f', '--font-path', help='Путь к шрифтам')
    parser.add_argument('-n', '--font-name', help='Имя шрифта (без расширения .ttf)', default='Montserrat')
    parser.add_argument('-b', '--background', help='Путь к фоновому изображению (поддерживаются png, jpg, jpeg)', default='bg')
    parser.add_argument('-t', '--text-color', help='Цвет текста в HEX-формате без #', default='1b1b1b')
    parser.add_argument('-c', '--heading-color', help='Цвет заголовков в HEX-формате без #', default='1b1b1b')
    parser.add_argument('--heading-gradient', help='Использовать градиент для заголовков (yes/no)', default='no')
    parser.add_argument('--heading-gradient-end', help='Конечный цвет градиента заголовков', default='1b1b1b')
    parser.add_argument('--text-gradient', help='Использовать градиент для текста (yes/no)', default='no')
    parser.add_argument('--text-gradient-end', help='Конечный цвет градиента текста', default='1b1b1b')
    parser.add_argument('--bg-effect', help='Эффект для фона (lighten/darken/none)', default='none')
    parser.add_argument('--bg-effect-intensity', help='Интенсивность эффекта (1-10)', default='0')
    
    args = parser.parse_args()
    
    input_file = args.input_file
    output_file = args.output if args.output else input_file.rsplit('.', 1)[0] + '.pdf'
    
    converter = MarkdownSlideConverter(
        font_path=args.font_path, 
        font_name=args.font_name,
        background_image=args.background,
        heading_color=args.heading_color,
        text_color=args.text_color,
        heading_gradient=args.heading_gradient,
        heading_gradient_end=args.heading_gradient_end,
        text_gradient=args.text_gradient,
        text_gradient_end=args.text_gradient_end,
        bg_effect=args.bg_effect,
        bg_effect_intensity=args.bg_effect_intensity
    )
    slides = converter.parse_markdown(input_file)
    converter.create_pdf(slides, output_file)

if __name__ == '__main__':
    main() 